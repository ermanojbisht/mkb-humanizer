"""Run with QT_QPA_PLATFORM=offscreen python -m unittest discover -s tests."""
import json
import os
from pathlib import Path
import tempfile
import threading
import time
import unittest
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from unittest.mock import patch

from PySide6.QtWidgets import QApplication, QMessageBox
from PySide6.QtCore import QSettings
from humanizer.gui import Window

DRAFT = ('This innovative solution serves as a testament to our commitment to excellence. '
         'Our team delivered 12 workshops in 2025, bringing practical training to local schools. '
         'The sessions helped teachers develop useful classroom skills and share their experiences.')
REWRITE = ('We ran 12 workshops in 2025 to give local schools practical training. '
           'Teachers worked on classroom skills and shared what they had learned. '
           'Our team built the sessions around everyday teaching needs, with clear examples and time for discussion.')


class Handler(BaseHTTPRequestHandler):
    mode = 'ok'
    def log_message(self, *args):
        pass
    def send_json(self, status, data):
        raw = json.dumps(data).encode()
        self.send_response(status)
        self.send_header('Content-Type', 'application/json')
        self.end_headers()
        try:
            self.wfile.write(raw)
        except BrokenPipeError:
            pass
    def do_GET(self):
        self.send_json(200, {'models': [{'name': 'hf.co/test/humanizer:Q4_K_M'}, {'name': 'other:latest'}]})
    def do_POST(self):
        payload = json.loads(self.rfile.read(int(self.headers['Content-Length'])))
        assert payload['raw'] is True
        assert '### Rewritten:' in payload['prompt']
        if self.mode == 'slow':
            time.sleep(2)
        if self.mode == 'error':
            self.send_json(400, {'error': 'test generation failure'})
        else:
            self.send_json(200, {'response': REWRITE, 'done': True})


class GuiIntegration(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.config = tempfile.TemporaryDirectory()
        QSettings.setDefaultFormat(QSettings.Format.IniFormat)
        QSettings.setPath(QSettings.Format.IniFormat, QSettings.Scope.UserScope, cls.config.name)
        cls.app = QApplication.instance() or QApplication([])
        cls.server = ThreadingHTTPServer(('127.0.0.1', 0), Handler)
        threading.Thread(target=cls.server.serve_forever, daemon=True).start()
    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.server.server_close()
        cls.config.cleanup()
    def wait(self, condition, timeout=15):
        deadline = time.monotonic() + timeout
        while not condition():
            self.app.processEvents()
            if time.monotonic() > deadline:
                self.fail('UI operation timed out')
            time.sleep(.01)
        self.app.processEvents()
    def setUp(self):
        Handler.mode = 'ok'
        self.window = Window(QSettings(str(Path(self.config.name) / 'test.ini'), QSettings.Format.IniFormat))
        self.window.url.setText(f'http://127.0.0.1:{self.server.server_port}')
        self.wait(lambda: self.window.discovery is None and self.window.model.count() > 0)
        self.messages = []
        self.window.message = self.messages.append
        self.tmp = tempfile.TemporaryDirectory()
    def tearDown(self):
        self.window.close()
        self.window.deleteLater()
        self.app.processEvents()
        self.tmp.cleanup()
    def file_job(self, suffix='.md'):
        source = Path(self.tmp.name) / ('source' + suffix)
        target = Path(self.tmp.name) / ('result' + suffix)
        if suffix == '.docx':
            from docx import Document
            doc = Document()
            doc.add_heading('Workshop report', 1)
            doc.add_paragraph(DRAFT)
            doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Preserved table'
            doc.save(source)
        else:
            source.write_text('# Workshop report\n\n' + DRAFT, encoding='utf-8')
        self.window.tabs.setCurrentIndex(1)
        self.window.input_path.setText(str(source))
        self.window.output_path.setText(str(target))
        return source, target
    def test_text_and_report(self):
        self.assertEqual(self.window.model.count(), 1)
        self.window.original.setPlainText(DRAFT)
        self.window.start_rewrite()
        self.assertFalse(self.window.tabs.isEnabled())
        self.wait(lambda: self.window.process is None)
        self.assertEqual(self.window.result.toPlainText(), REWRITE)
        self.assertTrue(self.window.export.isEnabled())
        self.assertEqual(self.window.report['input'], 'pasted text')
        self.assertEqual(self.messages, [])
    def test_markdown_preserves_original(self):
        source, target = self.file_job()
        original = source.read_bytes()
        self.window.start_rewrite()
        self.wait(lambda: self.window.process is None)
        self.assertEqual(source.read_bytes(), original)
        self.assertIn('# Workshop report', target.read_text())
        self.assertIn(REWRITE, target.read_text())
        self.assertEqual(self.window.report['output'], str(target))
    def test_docx_preserves_structures(self):
        from docx import Document
        source, target = self.file_job('.docx')
        original = source.read_bytes()
        self.window.start_rewrite()
        self.wait(lambda: self.window.process is None)
        doc = Document(target)
        self.assertEqual(doc.paragraphs[0].text, 'Workshop report')
        self.assertEqual(doc.tables[0].cell(0, 0).text, 'Preserved table')
        self.assertEqual(source.read_bytes(), original)
    def test_cancel_leaves_existing_output(self):
        Handler.mode = 'slow'
        _, target = self.file_job()
        target.write_text('existing output')
        with patch.object(QMessageBox, 'question', return_value=QMessageBox.StandardButton.Yes):
            self.window.start_rewrite()
        self.wait(lambda: 'using the' in self.window.log.toPlainText())
        self.window.cancel_job()
        self.wait(lambda: self.window.process is None)
        self.assertEqual(target.read_text(), 'existing output')
        self.assertIsNone(self.window.report)
        self.assertIsNone(self.window.temp)
    def test_failure_leaves_existing_output(self):
        Handler.mode = 'error'
        _, target = self.file_job()
        target.write_text('existing output')
        with patch.object(QMessageBox, 'question', return_value=QMessageBox.StandardButton.Yes):
            self.window.start_rewrite()
        self.wait(lambda: self.window.process is None)
        self.assertEqual(target.read_text(), 'existing output')
        self.assertTrue(self.messages)
    def test_reject_input_alias_and_empty_text(self):
        self.window.start_rewrite()
        self.assertIn('Paste', self.messages[-1])
        source, target = self.file_job()
        target.symlink_to(source)
        self.window.start_rewrite()
        self.assertIn('different', self.messages[-1])
        self.assertIsNone(self.window.process)
    def test_unreachable_ollama(self):
        self.window.url.setText('http://127.0.0.1:1')
        self.window.refresh_models()
        self.wait(lambda: self.window.discovery is None)
        self.assertFalse(self.window.rewrite.isEnabled())
        self.assertIn('unreachable', self.window.connection_status.text())


if __name__ == '__main__':
    unittest.main()
